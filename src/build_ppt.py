# -*- coding: utf-8 -*-
"""
06 生成 5 分钟汇报用 PPT（python-pptx，16:9，12 页）
输出 report/题2_汇报提纲.pptx
运行：python src/build_ppt.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import ROOT, FIG, METRICS

M = json.loads(METRICS.read_text(encoding="utf-8"))
P, E, AL = M["prepare"], M["eda"], M["alarm"]
A, A2, B, C, D = (M["task_a_random_split"], M["task_a_no_variable_id"],
                  M["task_b_leave_one_device"], M["task_c_consistency"],
                  M["task_d_early_warning"])
STR, DS = M["task_a_time_episode"], M["task_d_strict"]

OUT = ROOT / "report" / "题2_汇报提纲.pptx"
W, H = Inches(13.333), Inches(7.5)

NAVY = RGBColor(0x1F, 0x3B, 0x57)
ACCENT = RGBColor(0xE8, 0x74, 0x3B)
GREY = RGBColor(0x55, 0x5F, 0x6B)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"


def tb(slide, x, y, w, h, text, size=16, bold=False, color=GREY, align=PP_ALIGN.LEFT,
       line_spacing=1.25, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT
    return box


def bullets(slide, x, y, w, h, items, size=15, color=GREY, gap=8):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.3
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = "▪ " + it
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = FONT
    return box


def new_slide(prs, title, subtitle=None, band=True):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if band:
        bar = s.shapes.add_shape(1, 0, 0, W, Inches(1.05))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        bar.shadow.inherit = False
    tb(s, Inches(0.55), Inches(0.18), Inches(10.5), Inches(0.7), title, size=25, bold=True,
       color=WHITE if band else NAVY)
    if subtitle:
        tb(s, Inches(0.6), Inches(0.66), Inches(11), Inches(0.36), subtitle, size=12,
           color=RGBColor(0xC8, 0xD6, 0xE4) if band else GREY)
    return s


def pic(slide, name, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = w
    if h:
        kw["height"] = h
    return slide.shapes.add_picture(str(FIG / name), x, y, **kw)


def footer(slide, idx, text="题目二　车辆物联网遥测数据分析与告警识别"):
    tb(slide, Inches(0.55), Inches(7.02), Inches(9), Inches(0.32), text, size=10,
       color=RGBColor(0x99, 0xA3, 0xAE))
    tb(slide, Inches(12.3), Inches(7.02), Inches(0.7), Inches(0.32), f"{idx}", size=10,
       color=RGBColor(0x99, 0xA3, 0xAE), align=PP_ALIGN.RIGHT)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # ---------------- 1 封面
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    bg.shadow.inherit = False
    tb(s, Inches(1.0), Inches(1.55), Inches(11), Inches(0.5),
       "《Python 数据分析编程》课程项目汇报", size=16, color=RGBColor(0xBF, 0xCF, 0xE0))
    tb(s, Inches(1.0), Inches(2.25), Inches(11.3), Inches(1.0),
       "车辆物联网遥测数据分析与告警识别", size=40, bold=True, color=WHITE)
    line = s.shapes.add_shape(1, Inches(1.05), Inches(3.45), Inches(2.2), Emu(38100))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()
    tb(s, Inches(1.0), Inches(3.75), Inches(11), Inches(1.2),
       "173,853 条车载遥测 → 清洗与宽表 → 运行画像 → 告警语义 → 建模评估", size=18,
       color=RGBColor(0xD5, 0xDF, 0xEA))
    tb(s, Inches(1.0), Inches(5.5), Inches(11), Inches(0.9),
       "题目二　|　第 19 组　|　成员：周泽林　郭德宇　钟子恒", size=14,
       color=RGBColor(0x9F, 0xB3, 0xC8))

    # ---------------- 2 问题与数据
    s = new_slide(prs, "一、问题与数据", "长表 + 多频率 + 事件驱动：车联网数据的三个“难”")
    bullets(s, Inches(0.6), Inches(1.45), Inches(6.3), Inches(5.0), [
        f"规模：{P['raw_rows']:,} 行 6 列，3 台设备、{P['n_variables']} 种变量，7.5 天",
        "结构：长表（每行一个“变量—取值”），定位为“纬度,经度,海拔”字符串",
        f"稀疏：按(设备,时间戳)透视后填充率仅 {P['wide_ts_fill_rate']*100:.1f}%",
        "标签：alarmClass 附着在单条读数上（0 正常，1–5 告警）",
        "目标：清洗规整 → 还原运行状态 → 解释告警成因 → 建模识别与评估",
    ], size=15)
    pic(s, "fig01_数据概览.png", Inches(7.0), Inches(1.35), w=Inches(5.9))
    footer(s, 2)

    # ---------------- 3 技术路线
    s = new_slide(prs, "二、技术路线", "一步一脚本，指标全部落盘，可一键复现")
    steps = [("原始长表\n173,853 行", 0.6), ("清洗规整\n拆分/置空/类型", 3.05),
             ("长表转宽表\n时间戳·1min·5min", 5.5), ("探索分析\n画像·相关·时序", 7.95),
             ("告警识别\n片段·语义·阈值", 10.4)]
    for txt, x in steps:
        box = s.shapes.add_shape(5, Inches(x), Inches(1.9), Inches(2.15), Inches(1.25))
        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT
        box.line.color.rgb = RGBColor(0xC8, 0xD2, 0xDD)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        for i, ln in enumerate(txt.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = ln
            r.font.size = Pt(14 if i == 0 else 11)
            r.font.bold = i == 0
            r.font.color.rgb = NAVY if i == 0 else GREY
            r.font.name = FONT
    for x in [2.75, 5.2, 7.65, 10.1]:
        ar = s.shapes.add_shape(13, Inches(x), Inches(2.35), Inches(0.28), Inches(0.3))
        ar.fill.solid()
        ar.fill.fore_color.rgb = ACCENT
        ar.line.fill.background()
    bullets(s, Inches(0.6), Inches(3.6), Inches(12.2), Inches(3.0), [
        "四组实验：A 读数级多分类（A2 去变量身份）/ B 留一设备泛化 / C 无监督一致性 / D 提前预警",
        "环境：Python 3.14 + pandas 3.0 + scikit-learn 1.9，全流程约 1 分钟，指标写入 output/metrics.json",
        "报告 39 页、18 图、18 表；正文数字均由 metrics.json 读取，保证可复现一致",
    ], size=15)
    footer(s, 3)

    # ---------------- 4 清洗
    s = new_slide(prs, "三、数据清洗：把“脏且散”变成可用",
                  "原则：可疑值置空而非整行删除，保留同刻其他信号")
    bullets(s, Inches(0.6), Inches(1.4), Inches(6.2), Inches(5.2), [
        f"timeMili 缺失 {P['null_timeMili']:,}（70%）→ 由 timestamp 重建，两列一致",
        f"定位 {P['position_rows']:,} 条按逗号拆三列；{P['position_invalid']:,} 条越界置空（11.6%）",
        f"删除错标变量 PIDNameNotAvailable（{P['non_numeric_rows']} 条，取值为定位串）",
        "物理范围校验：转速 >8000、模块电压 <8V 等 104 条置空",
        f"清洗后有效读数 {P['clean_long_rows']:,} 条、定位 {P['clean_pos_rows']:,} 条",
        "关键发现：车-NAB0gBAQ 定位 81.4% 为 0 → 该车不做轨迹/里程分析",
    ], size=14)
    pic(s, "fig03_轨迹.png", Inches(6.95), Inches(1.9), w=Inches(5.95))
    tb(s, Inches(6.95), Inches(5.55), Inches(6.0), Inches(0.9),
       "轨迹呈“同一地点反复环绕”的短途特征（数据集主体位于印度卡纳塔克邦一带），\n"
       "对应车-twB1gBAQ 与车-GAB0gBAQ 的活动范围。", size=11)
    footer(s, 4)

    # ---------------- 5 宽表
    s = new_slide(prs, "四、长表转宽表：两种口径", "先对齐、再聚合")
    bullets(s, Inches(0.6), Inches(1.4), Inches(6.4), Inches(5.1), [
        f"时间戳级：{P['wide_ts_rows']:,} 行 × {P['wide_ts_cols']} 列，填充率 {P['wide_ts_fill_rate']*100:.1f}%",
        f"连续信号同设备前向填充（≤12 步）→ 可用率提升到 {P['wide_ts_fill_rate_ffill']*100:.1f}%",
        f"周期聚合：1 min {P['wide_1min']['rows']:,} 个窗口 / 5 min {P['wide_5min']['rows']:,} 个窗口",
        "每窗口附：均值·极值·标准差、GPS 里程与海拔、告警标签（各类别是否出现）、小时/星期",
        f"注意：5 min 窗口中也只有 {P['wide_5min']['speed_cover']*100:.1f}% 含车速读数——采样稀疏的真实代价",
        "→ 建模主战场放在“读数级”，窗口级仅用于趋势与预警实验",
    ], size=14)
    pic(s, "fig04_车速转速时序.png", Inches(7.1), Inches(1.45), w=Inches(5.8))
    tb(s, Inches(7.1), Inches(5.35), Inches(5.8), Inches(1.2),
       "2020-08-17：车速读数只出现在 10—11 时、14—17 时与 19 时，其余时间无 OBD 记录\n"
       "（事件驱动上报，熄火停放时仅定位工作）", size=11)
    footer(s, 5)

    # ---------------- 6 设备画像
    s = new_slide(prs, "五、三台车，三种用车画像", "结论必须分设备讨论")
    pic(s, "fig06_设备特征对比.png", Inches(0.45), Inches(1.25), w=Inches(7.55))
    x = Inches(8.25)
    tb(s, x, Inches(1.4), Inches(4.6), Inches(0.4), "车-GAB0gBAQ　高速长途", size=14, bold=True,
       color=ACCENT)
    tb(s, x, Inches(1.75), Inches(4.6), Inches(1.0),
       f"均速 {E['device_speed_mean']['GAB0gBAQ']} km/h，高速占 "
       f"{E['state_share']['高速(>60)']['GAB0gBAQ']*100:.1f}%，日均 {E['daily_km']['mean']['GAB0gBAQ']} km",
       size=12)
    tb(s, x, Inches(2.75), Inches(4.6), Inches(0.4), "车-NAB0gBAQ　城市起停", size=14, bold=True,
       color=ACCENT)
    tb(s, x, Inches(3.1), Inches(4.6), Inches(1.0),
       f"怠速 {E['state_share']['怠速']['NAB0gBAQ']*100:.1f}%、低速 "
       f"{E['state_share']['低速(<30)']['NAB0gBAQ']*100:.1f}%，均速 {E['device_speed_mean']['NAB0gBAQ']} km/h",
       size=12)
    tb(s, x, Inches(4.1), Inches(4.6), Inches(0.4), "车-twB1gBAQ　短途停放", size=14, bold=True,
       color=ACCENT)
    tb(s, x, Inches(4.45), Inches(4.6), Inches(1.0),
       f"熄火/静止占 {E['state_share']['熄火/静止']['twB1gBAQ']*100:.1f}%，"
       f"均速 {E['device_speed_mean']['twB1gBAQ']} km/h，日均 {E['daily_km']['mean']['twB1gBAQ']} km",
       size=12)
    tb(s, x, Inches(5.45), Inches(4.6), Inches(1.0),
       "车速—转速 Spearman 相关 0.74；\n纵向与垂向加速度弱负相关（加速抬头/减速点头）",
       size=11.5, color=RGBColor(0x77, 0x80, 0x8A))
    footer(s, 6)

    # ---------------- 7 告警分布与语义
    s = new_slide(prs, "六、告警是什么？", f"48.3% 读数带告警 → 实为“分级驾驶行为/工况”标签")
    pic(s, "fig09_告警信号分布.png", Inches(0.35), Inches(1.05), w=Inches(6.6))
    x = Inches(8.5)
    rows = [("C1", "低速强冲击", "车速 7.6｜横向 −8.7"), ("C2", "中低速工况", "车速 23.9｜节气门通道"),
            ("C3", "中高速加速", "车速 49.8｜纵向 +1.0"), ("C4", "高速高负荷", "车速 62.1｜转速 2324"),
            ("C5", "熄火静止", "IGNITION/IDLING=0")]
    for i, (c, name, ev) in enumerate(rows):
        y = Inches(1.35 + i * 0.95)
        tb(s, x, y, Inches(0.7), Inches(0.4), c, size=15, bold=True, color=ACCENT)
        tb(s, x + Inches(0.7), y, Inches(4.0), Inches(0.4), name, size=14, bold=True, color=NAVY)
        tb(s, x + Inches(0.7), y + Inches(0.33), Inches(4.1), Inches(0.4), ev, size=11)
    tb(s, x, Inches(6.35), Inches(4.4), Inches(1.0),
       "阈值几乎可“读出”：车速<20→C1，20–30→C2，50–60→C3/C4，>80→C4", size=11.5,
       color=RGBColor(0x77, 0x80, 0x8A))
    footer(s, 7)

    # ---------------- 8 告警与状态 + 片段
    s = new_slide(prs, "七、告警与运行状态、事件片段", "1,078 段告警片段，高等级集中在午后高峰")
    pic(s, "fig11_告警与运行状态.png", Inches(0.4), Inches(1.3), w=Inches(7.6))
    pic(s, "fig13_告警片段.png", Inches(8.15), Inches(1.3), w=Inches(4.6))
    c1s = AL["before_after"]["by_class"]["1"]["speed"]["delta"]
    c3s = AL["before_after"]["by_class"]["3"]["speed"]["delta"]
    c3r = AL["before_after"]["by_class"]["3"]["rpm"]["delta"]
    tb(s, Inches(0.45), Inches(4.75), Inches(7.5), Inches(1.6),
       "车速 >50 km/h 的读数 100% 带告警（C3 22.3% / C4 77.7%）；\n"
       f"告警读数相对前 5 分钟：C3 车速 +{c3s:.0f}、转速 +{c3r:.0f}；C1 车速 {c1s:.0f}；\n"
       "各片段中位时长 0–5 分钟，C4 片段集中在 14–19 时 → 高风险时段可预判", size=12)
    tb(s, Inches(8.15), Inches(4.62), Inches(4.8), Inches(1.6),
       "C3 片段最多（491 段），C5 最长（p90 27.9 分钟）；\n"
       "车-twB1gBAQ 931 段，14 时最多", size=12)
    footer(s, 8)

    # ---------------- 9 建模
    s = new_slide(prs, "八、建模：能识别，但不等于“懂车”",
                  "读数级 6 分类 + 三组检验")
    rows = [["模型", "准确率", "宏 F1"],
            ["多数类基线", f"{A['基线(多数类)']['acc']:.3f}", f"{A['基线(多数类)']['macro_f1']:.3f}"],
            ["逻辑回归", f"{A['逻辑回归']['acc']:.3f}", f"{A['逻辑回归']['macro_f1']:.3f}"],
            ["随机森林", f"{A['随机森林']['acc']:.3f}", f"{A['随机森林']['macro_f1']:.3f}"],
            ["梯度提升", f"{A['直方图梯度提升']['acc']:.3f}", f"{A['直方图梯度提升']['macro_f1']:.3f}"]]
    tw = Inches(4.1)
    tbl = s.shapes.add_table(len(rows), 3, Inches(0.55), Inches(1.4), tw, Inches(2.4)).table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = v
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(12 if i else 12.5)
                r.font.bold = i == 0
                r.font.name = FONT
                r.font.color.rgb = WHITE if i == 0 else GREY
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (
                LIGHT if i % 2 == 0 else WHITE)
    bullets(s, Inches(0.55), Inches(4.15), Inches(4.6), Inches(2.7), [
        f"去掉变量身份 → 随机森林仍 {A2['随机森林(无变量身份)']['acc']:.3f}、逻辑回归仅 {A2['逻辑回归(无变量身份)']['acc']:.3f}",
        "留一设备 → 宏 F1 降到 0.53–0.66",
        f"时间+事件隔离 → 宏 F1 {STR['models']['直方图梯度提升']['macro_f1']:.3f}",
        "结论：高精度来自“复现规则通道”",
    ], size=11.5)
    pic(s, "fig15_混淆矩阵.png", Inches(5.2), Inches(1.35), w=Inches(7.7))
    footer(s, 9)

    # ---------------- 10 一致性与预警
    s = new_slide(prs, "九、两个“反直觉”的检验", "统计异常 ≠ 业务告警；预警只有弱信号")
    pic(s, "fig17_异常一致性.png", Inches(0.4), Inches(1.3), w=Inches(6.3))
    pic(s, "fig18_提前预警.png", Inches(6.85), Inches(1.3), w=Inches(6.1))
    tb(s, Inches(0.45), Inches(5.25), Inches(6.2), Inches(1.6),
       "3σ 异常点对告警的精确率 0–100%（多数变量 ≥90%），召回率仅 0.2%–7.7%\n"
       "→ 极端值几乎必被规则捕获，但业务告警多来自“正常范围内的越阈值”", size=12)
    tb(s, Inches(6.9), Inches(5.15), Inches(6.2), Inches(1.9),
       "过去 10 min → 未来 10 min 是否出现 C4 片段：\n"
       f"逻辑回归 AUC {D['逻辑回归']['auc']:.3f}（严格 {DS['逻辑回归']['auc']:.3f}）、F1 {D['逻辑回归']['f1']:.3f}\n"
       "→ 弱可用：宜作风险评分，与规则串联使用", size=12)
    footer(s, 10)

    # ---------------- 11 结论与建议
    s = new_slide(prs, "十、结论与建议", "给业务可落地的四件事")
    bullets(s, Inches(0.6), Inches(1.35), Inches(6.2), Inches(5.4), [
        "数据：事件驱动上报导致稀疏（时间戳级填充率 9%），定位脏数据 11.6%",
        "画像：高速长途 / 城市起停 / 短途停放三类，告警率 30%–58% 差异显著",
        f"告警：分级行为标签，{AL['episodes']} 段事件；等级与车速、转速、加速度正相关",
        f"建模：读数级宏 F1 {A['直方图梯度提升']['macro_f1']:.3f}，跨设备 0.53–0.66；预警 AUC {D['逻辑回归']['auc']}",
    ], size=14)
    box = s.shapes.add_shape(5, Inches(7.0), Inches(1.4), Inches(5.8), Inches(5.1))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.color.rgb = RGBColor(0xD3, 0xDB, 0xE4)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.35)
    tf.margin_top = Inches(0.3)
    items = ["① 告警分层：C5 熄火改“状态标记”，释放 6.9% 告警通道",
             "② 提速采样：车速/转速/加速度改 1–5 秒定频上报",
             "③ 跨车建模先对齐：新车先跑 1–2 天“只评估不告警”",
             "④ 驾驶风险分：片段数 × 等级权重 ÷ 里程，优先管控高风险车"]
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.35
        p.space_after = Pt(16)
        r = p.add_run()
        r.text = it
        r.font.size = Pt(14)
        r.font.color.rgb = NAVY
        r.font.name = FONT
    footer(s, 11)

    # ---------------- 12 分工
    s = new_slide(prs, "十一、分工与致谢", "模块负责 + 交叉复核，Git 分支与分工一一对应")
    rows = [["成员", "负责模块", "对应分支"],
            ["周泽林（组长）", "数据工程与清洗（prepare / common）、报告第 2–3 章", "feature/data-pipeline"],
            ["郭德宇", "探索与告警分析（eda / alarm）、报告第 4–5 章", "feature/eda-alarm"],
            ["钟子恒", "建模评估与报告排版（model / build_report）、汇报视频", "feature/modeling-report"]]
    tbl = s.shapes.add_table(4, 3, Inches(0.6), Inches(1.5), Inches(12.1), Inches(2.0)).table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = v
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j else PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(12.5)
                r.font.bold = i == 0
                r.font.name = FONT
                r.font.color.rgb = WHITE if i == 0 else GREY
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
    tb(s, Inches(0.6), Inches(3.95), Inches(12), Inches(1.6),
       "复现方式：python run_all.py（约 1 分钟）　|　"
       "仓库：https://github.com/337437745-collab/vehicle_analysis\n"
       "报告：report/题2_车辆物联网遥测数据分析与告警识别_报告.docx（39 页 / 18 图 / 18 表）",
       size=13)
    tb(s, Inches(0.6), Inches(5.6), Inches(12), Inches(0.8), "谢谢各位老师与同学，欢迎提问！", size=22,
       bold=True, color=ACCENT)
    footer(s, 12)

    prs.save(OUT)
    print(f"PPT 已生成：{OUT}（{len(prs.slides.__iter__.__self__._sldIdLst)} 页）")
    return OUT


if __name__ == "__main__":
    main()


